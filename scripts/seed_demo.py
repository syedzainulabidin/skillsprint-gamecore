"""
Seed a demo fictional company: 10 job roles, sample documents, requirements,
and role-requirement mappings so the app has a demo dataset without needing
to click through the UI.

Run from the backend directory:
    venv\\Scripts\\python.exe scripts\\seed_demo.py
"""

import io
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from datetime import date

from app.core.constants import (
    DUE_STAGES,
    MUST_TYPES,
    PRIORITIES,
    REQUIREMENT_TYPES,
)
from app.database.connection import Base, SessionLocal, engine
from app.models import (  # noqa: F401
    Document,
    DocumentVersion,
    JobRole,
    Requirement,
    RoleRequirement,
    User,
)
from app.services import (
    document_service,
    job_role_service,
    prompt_service,
    precedence_rule_service,
    requirement_service,
    role_requirement_service,
    user_service,
)
from app.schemas.job_role import JobRoleCreate
from app.schemas.requirement import RequirementCreate
from app.schemas.role_requirement import RoleRequirementCreate
from app.schemas.user import UserCreate


ROLES = [
    ("Sales Executive", "Sales"),
    ("Customer Support Executive", "Support"),
    ("HR Executive", "Human Resources"),
    ("Finance Associate", "Finance"),
    ("Operations Coordinator", "Operations"),
    ("Marketing Executive", "Marketing"),
    ("Software Support Engineer", "Engineering"),
    ("Branch Manager", "Sales"),
    ("Data Analyst", "Analytics"),
    ("Team Leader", "Operations"),
]

DOCS = [
    ("POL-01", "Company Handbook", "handbook", None, "General onboarding for all employees. Working hours, dress code, communication."),
    ("POL-02", "HR Policy", "hr_policy", "Human Resources", "Leave, benefits, performance reviews, code of conduct."),
    ("POL-03", "Data Privacy Policy", "data_privacy", None, "Customer data handling, consent, retention, deletion."),
    ("POL-04", "Information Security Policy", "info_security", None, "Passwords, MFA, phishing, incident reporting."),
    ("POL-05", "Workplace Conduct Policy", "workplace_conduct", None, "Anti-harassment, respect, escalation."),
    ("SOP-01", "Sales Escalation SOP", "sop", "Sales", "How to escalate customer complaints within 24 hours."),
    ("SOP-02", "Customer Support Ticket SOP", "sop", "Support", "Ticket triage, priority levels, SLA."),
    ("SOP-03", "Payroll Processing SOP", "sop", "Finance", "Payroll cycle, tax deductions, approvals."),
    ("SOP-04", "Marketing Campaign SOP", "sop", "Marketing", "Campaign approval workflow, brand review."),
    ("SOP-05", "Incident Response SOP", "sop", "Engineering", "Detect, contain, eradicate, recover from incidents."),
    ("FAQ-01", "General FAQ", "faq", None, "Common employee questions."),
    ("FAQ-02", "Benefits FAQ", "faq", "Human Resources", "Health insurance, retirement, PTO."),
    ("COMP-01", "GDPR Compliance", "compliance", None, "European data-protection requirements."),
    ("COMP-02", "SOC 2 Controls", "compliance", "Engineering", "Access control, change management, monitoring."),
    ("SAFE-01", "Workplace Safety", "safety", None, "Fire, first aid, evacuation."),
    ("PROC-01", "Employee Onboarding Process", "process_manual", "Human Resources", "Day 1 through 90 days."),
    ("PROC-02", "Change Management Process", "process_manual", "Engineering", "How code changes reach production."),
    ("ROLE-01", "Sales Executive Role Description", "role_description", "Sales", "Sales quotas, pipeline management, reporting."),
    ("ROLE-02", "Software Support Engineer Role Description", "role_description", "Engineering", "On-call rotation, triage, RCA."),
    ("DEPT-01", "Finance Department Guidelines", "department_guideline", "Finance", "Expense claims, purchase orders, month-end close."),
]


def _make_docx_bytes(title: str, sections: list) -> bytes:
    from docx import Document as DocxDocument

    buf = io.BytesIO()
    doc = DocxDocument()
    doc.add_heading(title, level=0)
    for i, (heading, body) in enumerate(sections, start=1):
        doc.add_heading(f"{i} {heading}", level=1)
        doc.add_paragraph(body)
    doc.save(buf)
    return buf.getvalue()


REQUIREMENTS = [
    ("R001", "Understand data privacy basics", "policy", "must_know", "high", "week_1", "POL-03", "1"),
    ("R002", "Complete information security training", "policy", "must_complete", "high", "week_1", "POL-04", "1"),
    ("R003", "Read and acknowledge workplace conduct policy", "policy", "must_acknowledge", "medium", "day_1", "POL-05", "1"),
    ("R004", "Understand payroll and benefits", "knowledge", "must_know", "medium", "week_2", "POL-02", "1"),
    ("R005", "Complete GDPR compliance module", "policy", "must_complete", "high", "first_30_days", "COMP-01", "1"),
    ("R006", "Sales pipeline management competency", "competency", "must_demonstrate", "high", "first_30_days", "ROLE-01", "1"),
    ("R007", "Follow sales escalation procedure", "process", "must_complete", "high", "week_2", "SOP-01", "1"),
    ("R008", "Customer support ticket handling", "process", "must_demonstrate", "high", "week_2", "SOP-02", "1"),
    ("R009", "Incident response procedure", "process", "must_complete", "critical", "first_30_days", "SOP-05", "1"),
    ("R010", "SOC 2 controls awareness", "policy", "must_know", "high", "first_60_days", "COMP-02", "1"),
    ("R011", "Payroll processing workflow", "process", "must_demonstrate", "high", "first_30_days", "SOP-03", "1"),
    ("R012", "Marketing campaign approval process", "process", "must_complete", "medium", "first_30_days", "SOP-04", "1"),
    ("R013", "Workplace safety training", "policy", "must_complete", "high", "week_1", "SAFE-01", "1"),
    ("R014", "Complete employee onboarding checklist", "task", "must_complete", "high", "first_30_days", "PROC-01", "1"),
    ("R015", "Understand change management", "knowledge", "must_know", "medium", "first_30_days", "PROC-02", "1"),
]


ROLE_REQUIREMENTS = {
    "Sales Executive": ["R001", "R002", "R003", "R006", "R007", "R013", "R014"],
    "Customer Support Executive": ["R001", "R002", "R003", "R005", "R008", "R013", "R014"],
    "HR Executive": ["R001", "R002", "R003", "R004", "R013", "R014"],
    "Finance Associate": ["R001", "R002", "R003", "R011", "R013", "R014"],
    "Operations Coordinator": ["R001", "R002", "R003", "R013", "R014"],
    "Marketing Executive": ["R001", "R002", "R003", "R012", "R013", "R014"],
    "Software Support Engineer": ["R001", "R002", "R003", "R009", "R010", "R013", "R014", "R015"],
    "Branch Manager": ["R001", "R002", "R003", "R006", "R007", "R013", "R014"],
    "Data Analyst": ["R001", "R002", "R003", "R005", "R010", "R013", "R014"],
    "Team Leader": ["R001", "R002", "R003", "R009", "R013", "R014"],
}


def main() -> int:
    Base.metadata.create_all(bind=engine)
    db = SessionLocal()

    try:
        precedence_rule_service.ensure_defaults(db)
        prompt_service.ensure_default_prompts(db)
        db.commit()

        admin = db.query(User).filter(User.system_role == "admin").first()
        if admin is None:
            print("No admin user; run the server once first (bootstrap creates admin).")
            return 1
        actor_id = admin.id
        print(f"Using admin #{actor_id} as author.")

        role_by_name = {}
        for name, dept in ROLES:
            existing = db.query(JobRole).filter(JobRole.name == name).first()
            if existing:
                role_by_name[name] = existing
                continue
            role = job_role_service.create_job_role(
                db,
                JobRoleCreate(
                    name=name,
                    department=dept,
                    description=f"Onboarding role: {name}",
                ),
            )
            role_by_name[name] = role
            db.commit()
        print(f"Roles ready: {len(role_by_name)}")

        doc_by_code = {}
        for doc_code, name, doc_type, department, desc in DOCS:
            existing = db.query(Document).filter(Document.doc_code == doc_code).first()
            if existing:
                doc_by_code[doc_code] = existing
                continue
            sections = [
                ("Overview", desc),
                ("Scope", f"This applies to all employees in {department or 'the company'}."),
                ("Key requirements", "Follow the policy at all times. Ask HR if unsure."),
                ("Compliance", "Non-compliance may result in disciplinary action."),
            ]
            data = _make_docx_bytes(name, sections)
            document, version = document_service.upload_document(
                db,
                doc_code=doc_code,
                name=name,
                doc_type=doc_type,
                department=department,
                category=None,
                description=desc,
                version_label="v1",
                effective_date=date.today(),
                expiry_date=None,
                filename=f"{doc_code.lower()}.docx",
                mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                data=data,
                uploaded_by=actor_id,
            )
            doc_by_code[doc_code] = document
            db.commit()
        print(f"Documents ready: {len(doc_by_code)}")

        req_by_code = {}
        for code, title, rtype, mtype, prio, stage, doc_code, section in REQUIREMENTS:
            existing = db.query(Requirement).filter(Requirement.req_code == code).first()
            if existing:
                req_by_code[code] = existing
                continue
            doc = doc_by_code.get(doc_code)
            req = requirement_service.create_requirement(
                db,
                RequirementCreate(
                    req_code=code,
                    title=title,
                    description=f"Auto-seeded requirement: {title}",
                    requirement_type=rtype,
                    must_type=mtype,
                    priority=prio,
                    due_stage=stage,
                    competency=None,
                    assessment_topic=None,
                    source_document_id=doc.id if doc else None,
                    source_section=section,
                    source_version_id=None,
                    source_chunk_id=None,
                ),
                actor_id,
            )
            req_by_code[code] = req
            db.commit()
        print(f"Requirements ready: {len(req_by_code)}")

        assigned = 0
        for role_name, req_codes in ROLE_REQUIREMENTS.items():
            role = role_by_name[role_name]
            for req_code in req_codes:
                req = req_by_code.get(req_code)
                if req is None:
                    continue
                existing = (
                    db.query(RoleRequirement)
                    .filter(
                        RoleRequirement.job_role_id == role.id,
                        RoleRequirement.requirement_id == req.id,
                    )
                    .first()
                )
                if existing:
                    continue
                role_requirement_service.assign(
                    db,
                    RoleRequirementCreate(
                        job_role_id=role.id,
                        requirement_id=req.id,
                        is_mandatory=req.must_type in {
                            "must_know",
                            "must_complete",
                            "must_demonstrate",
                            "must_acknowledge",
                        },
                    ),
                    actor_id,
                )
                assigned += 1
        db.commit()
        print(f"Role-requirement assignments created: {assigned}")

        demo_employees = [
            ("EMP001", "Aiden Sales", "aiden@example.com", "Sales Executive", "Sales"),
            ("EMP002", "Priya Support", "priya@example.com", "Customer Support Executive", "Support"),
            ("EMP003", "Rahul Finance", "rahul@example.com", "Finance Associate", "Finance"),
        ]
        created = 0
        for emp_id, name, email, role_name, dept in demo_employees:
            if db.query(User).filter(User.email == email).first():
                continue
            role = role_by_name[role_name]
            user_service.create_user(
                db,
                UserCreate(
                    employee_id=emp_id,
                    name=name,
                    email=email,
                    password="password123",
                    system_role="employee",
                    job_role_id=role.id,
                    department=dept,
                    experience_level="beginner",
                    joining_date=date.today(),
                ),
            )
            created += 1
        db.commit()
        print(f"Demo employees created: {created}")

        print()
        print("Seed complete.")
        print("Default admin login: admin@gmail.com / 123456789")
        print("Demo employee login: aiden@example.com / password123")
        return 0
    finally:
        db.close()


if __name__ == "__main__":
    sys.exit(main())
